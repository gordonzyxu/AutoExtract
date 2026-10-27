"""
API tests for /api/extract and /api/export (Flask test client, no server).

    .venv/Scripts/python -m pytest tests -q
"""

from __future__ import annotations

import io
import json
import os
import sys

import cv2
import numpy as np
import openpyxl
import pytest
from flask import Flask

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "backend"))
sys.path.insert(0, HERE)

import extract_routes  # noqa: E402
from batch_extract import iter_inputs  # noqa: E402
from pdf_utils import figures  # noqa: E402
from synth import BoxSpec, ChartSpec, LineSpec, ViolinSpec, render, render_box, render_line, render_violin  # noqa: E402

SPEC = ChartSpec([12.0, 31.5, 22.0], [1.5, 3.0, 2.2], labels=["Control", "CUMS", "CUMS+EA"])


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    extract_routes.JOBS_DIR = str(tmp_path_factory.mktemp("jobs"))
    app = Flask(__name__)
    app.add_url_rule("/api/extract", "extract", extract_routes.handle_extract, methods=["POST"])
    app.add_url_rule("/api/export", "export", extract_routes.handle_export, methods=["POST"])
    app.add_url_rule("/static/jobs/<job>/<filename>", "job_file", extract_routes.handle_job_file)
    return app.test_client()


def _png_bytes(spec=SPEC) -> bytes:
    ok, buf = cv2.imencode(".png", render(spec))
    assert ok
    return buf.tobytes()


def _pdf_bytes(caption: str) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(60, 60, 420, 400), stream=_png_bytes())
    page.insert_text((60, 430), caption, fontsize=9)
    return doc.tobytes()


def _pdf_chart_bytes(image) -> bytes:
    """Embed the rendered chart as a native-resolution figure in a PDF."""
    import pymupdf
    ok, buf = cv2.imencode(".png", image)
    assert ok
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(60, 60, 500, 500), stream=buf.tobytes())
    page.insert_text((60, 530), "Fig. 2. Experimental results.", fontsize=9)
    return doc.tobytes()


def _vector_pdf_bytes() -> bytes:
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(4, 3))
    ax.bar([0, 1, 2], [12, 25, 19], yerr=[1, 2, 2], capsize=3,
           facecolor="white", edgecolor="black")
    ax.set_ylim(0, 35)
    ax.set_xticks([0, 1, 2], ["A", "B", "C"])
    ax.set_ylabel("Value")
    fig.tight_layout()
    out = io.BytesIO()
    fig.savefig(out, format="pdf")
    plt.close(fig)
    return out.getvalue()


def _post(client, data: bytes, name: str, **form):
    payload = {"file": (io.BytesIO(data), name)}
    payload.update(form)
    return client.post("/api/extract", data=payload, content_type="multipart/form-data")


def test_extract_png(client):
    r = _post(client, _png_bytes(), "chart.png")
    assert r.status_code == 200, r.get_json()
    data = r.get_json()
    assert len(data["figures"]) == 1
    panel = data["figures"][0]["panels"][0]
    assert [round(b["mean"], 1) for b in panel["bars"]] == pytest.approx([12.0, 31.5, 22.0], abs=0.4)
    assert [b["label"] for b in panel["bars"]] == ["Control", "CUMS", "CUMS+EA"]
    assert panel["error_type"] == "sem"
    # image is served back
    img = client.get(data["figures"][0]["image_url"])
    assert img.status_code == 200 and len(img.data) > 1000


def test_extract_pdf_reads_caption(client):
    r = _post(client, _pdf_bytes("Fig. 3. Behavioural tests. Data are mean ± SD (n = 8)."), "paper.pdf")
    assert r.status_code == 200, r.get_json()
    fig = r.get_json()["figures"][0]
    assert fig["figure"] == "3"
    assert fig["error_type_from_caption"] == "sd"
    assert fig["panels"][0]["error_type"] == "sd"


@pytest.mark.parametrize("above", [True, False])
def test_pdf_reused_image_placements_keep_their_own_captions(tmp_path, above):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=900)
    png = _png_bytes()
    page.insert_text((50, 45 if above else 370), "Fig. 1. First chart. Mean SEM.", fontsize=9)
    page.insert_image(pymupdf.Rect(50, 60, 500, 350), stream=png)
    page.insert_text((50, 455 if above else 780), "Fig. 2. Second chart. Mean SD.", fontsize=9)
    page.insert_image(pymupdf.Rect(50, 470, 500, 760), stream=png)
    pdf = tmp_path / "reused.pdf"
    pdf.write_bytes(doc.tobytes())
    doc.close()

    found = figures(str(pdf))
    assert [f["figure"] for f in found] == ["1", "2"]
    assert [f["error_type"] for f in found] == ["sem", "sd"]


def test_pdf_keeps_smaller_full_figure_but_skips_page_logo(tmp_path):
    import pymupdf

    chart = cv2.resize(render(SPEC), (500, 400))
    ok, buf = cv2.imencode(".png", chart)
    assert ok
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(60, 70, 410, 350), stream=buf.tobytes())
    page.insert_text((60, 375), "Fig. 4. Small raster chart.", fontsize=9)
    page.insert_image(pymupdf.Rect(510, 20, 560, 60), stream=buf.tobytes())
    pdf = tmp_path / "small_figure.pdf"
    pdf.write_bytes(doc.tobytes())
    doc.close()

    found = figures(str(pdf))
    assert len(found) == 1
    assert found[0]["figure"] == "4"
    assert found[0]["image"].shape[:2] == (400, 500)


def test_pdf_renders_sparse_vector_axes_with_caption(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    shape = page.new_shape()
    shape.draw_line((90, 360), (450, 360))
    shape.draw_line((90, 360), (90, 100))
    shape.finish(color=(0, 0, 0), width=1)
    shape.commit()
    page.insert_text((90, 390), "Fig. 5. Vector chart.", fontsize=9)
    pdf = tmp_path / "sparse_vector.pdf"
    pdf.write_bytes(doc.tobytes())
    doc.close()

    found = figures(str(pdf))
    assert len(found) == 1
    assert found[0]["figure"] == "5"


def test_pdf_ruled_text_page_is_not_rendered_as_a_figure(tmp_path):
    import pymupdf

    # A title page has dozens of short rules and boxes but no axes; rendering
    # such pages at 300 dpi and scanning them only costs time.
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    for k in range(30):
        shape = page.new_shape()
        y = 80 + 22 * k
        shape.draw_line((60, y), (60 + 40 + 3 * k, y))
        shape.finish(color=(0, 0, 0), width=0.5)
        shape.commit()
    assert len(page.get_drawings()) >= 20
    page.insert_text((60, 60), "Fig. 1. Not a chart.", fontsize=9)
    pdf = tmp_path / "ruled.pdf"
    pdf.write_bytes(doc.tobytes())
    doc.close()

    assert figures(str(pdf)) == []


def test_pdf_side_by_side_figures_match_horizontal_captions(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=500)
    png = _png_bytes()
    page.insert_image(pymupdf.Rect(35, 60, 275, 300), stream=png)
    page.insert_image(pymupdf.Rect(320, 60, 560, 300), stream=png)
    page.insert_text((35, 330), "Fig. 6. Left chart.", fontsize=9)
    page.insert_text((320, 330), "Fig. 7. Right chart.", fontsize=9)
    pdf = tmp_path / "side_by_side.pdf"
    pdf.write_bytes(doc.tobytes())
    doc.close()

    assert [f["figure"] for f in figures(str(pdf))] == ["6", "7"]


def test_pdf_wrapped_caption_keeps_error_type(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=500)
    page.insert_image(pymupdf.Rect(50, 40, 450, 320), stream=_png_bytes())
    page.insert_textbox(pymupdf.Rect(50, 340, 450, 390),
                        "Fig. 8. Experimental outcomes.\nError bars show standard deviation.",
                        fontsize=9)
    pdf = tmp_path / "wrapped_caption.pdf"
    pdf.write_bytes(doc.tobytes())
    doc.close()

    found = figures(str(pdf))
    assert len(found) == 1
    assert found[0]["figure"] == "8" and found[0]["error_type"] == "sd"


def test_batch_overlapping_paths_are_processed_once(tmp_path):
    first = tmp_path / "a.png"
    second_dir = tmp_path / "sub"
    second_dir.mkdir()
    second = second_dir / "b.pdf"
    first.write_bytes(b"x")
    second.write_bytes(b"x")
    assert iter_inputs([str(first), str(tmp_path), str(second_dir), str(second)]) == [str(first), str(second)]


def test_invalid_extract_options_leave_no_job_directory(client):
    job = "invalidoptionsshouldnotcreateajob"
    path = os.path.join(extract_routes.JOBS_DIR, job)
    assert not os.path.exists(path)
    response = _post(client, _png_bytes(), "chart.png", job=job, options="{")
    assert response.status_code == 400
    assert not os.path.exists(path)


@pytest.mark.parametrize("kind,image,expected_groups", [
    ("bar", render(ChartSpec([12, 25, 19], [1, 2, 2])), 3),
    ("dot", render(ChartSpec([110, 170, 126, 116], [44, 52, 35, 45], kind="dot", dots=True,
                              colors=["#1f77b4", "#d62728", "#2ca02c", "#9467bd"], ylim=(0, 300))), 4),
    ("box", render_box(BoxSpec([(2, 5, 7, 9, 13), (4, 8, 11, 15, 19)], ylim=(0, 20))), 2),
    ("violin", render_violin(ViolinSpec([np.random.default_rng(1).normal(10, 2, 20),
                                           np.random.default_rng(2).normal(14, 2, 20)],
                                          inner="lines")), 2),
    ("line", render_line(LineSpec([0, 7, 14, 21], [[20, 22, 25, 27], [16, 17, 18, 19]],
                                  errs=[[1] * 4, [0.8] * 4])), 2),
])
def test_pdf_extracts_each_supported_chart_type(client, kind, image, expected_groups):
    r = _post(client, _pdf_chart_bytes(image), "figure.pdf",
              options=json.dumps({"chart_type": kind, "bar_names": False, "titles": False}))
    assert r.status_code == 200, r.get_json()
    figures = r.get_json()["figures"]
    assert len(figures) == 1 and len(figures[0]["panels"]) == 1
    panel = figures[0]["panels"][0]
    assert panel["kind"] == kind
    groups = panel["series"] if kind == "line" else panel["boxes"] if kind in ("box", "violin") else panel["bars"]
    assert len(groups) == expected_groups


def test_vector_pdf_with_few_paths_is_extracted(client):
    r = _post(client, _vector_pdf_bytes(), "vector.pdf",
              options=json.dumps({"chart_type": "bar", "bar_names": False, "titles": False}))
    assert r.status_code == 200, r.get_json()
    figures = r.get_json()["figures"]
    assert len(figures) == 1
    bars = figures[0]["panels"][0]["bars"]
    assert len(bars) == 3
    assert [b["mean"] for b in bars] == pytest.approx([12, 25, 19], abs=0.5)


def test_error_type_override_and_options(client):
    r = _post(client, _png_bytes(), "chart.png", error_type="sd",
              options=json.dumps({"bar_names": False, "titles": False}))
    panel = r.get_json()["figures"][0]["panels"][0]
    assert panel["error_type"] == "sd"
    assert all(b["label"] == "" for b in panel["bars"])
    assert panel["y_title"] == ""


@pytest.mark.parametrize("form,name,data,status", [
    ({}, None, None, 400),                                     # no file
    ({}, "notes.txt", b"hello", 400),                          # unsupported type
    ({"options": "{not json"}, "chart.png", None, 400),        # bad options
    ({"error_type": "variance"}, "chart.png", None, 400),      # bad error type
    ({"job": "../../etc"}, "chart.png", None, 400),            # path tricks in job id
])
def test_extract_rejects_bad_input(client, form, name, data, status):
    if name is None:
        r = client.post("/api/extract", data={}, content_type="multipart/form-data")
    else:
        r = _post(client, data if data is not None else _png_bytes(), name, **form)
    assert r.status_code == status
    assert "error" in r.get_json()


def test_extract_corrupt_files(client):
    r = _post(client, b"%PDF-1.4 garbage", "broken.pdf")
    assert r.status_code in (400, 500) and "error" in r.get_json()
    r = _post(client, b"\x89PNG not really", "broken.png")
    assert r.status_code == 200 and r.get_json()["figures"] == []


def test_export_roundtrip(client):
    data = _post(client, _png_bytes(), "chart.png").get_json()
    fig = data["figures"][0]
    panel = fig["panels"][0]
    panel["error_type"] = "sd"
    for b in panel["bars"]:
        b["n"] = 8
    panel["bars"][0]["label"] = "control"
    r = client.post("/api/export", json={"figures": [fig], "filename": "x.xlsx"})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[:2] == ["Summary", "All data"]
    ws = wb[wb.sheetnames[2]]
    header = [c.value for c in ws[1]][:12]
    assert header[0] == "mean" and header[-1] == "mean+sd" and header[5] == "treatment_group"
    row = [c.value for c in ws[2]][:12]
    assert row[5] == "control" and row[6] == 8
    assert abs(row[0] - 12.0) < 0.4
    assert isinstance(row[1], float) and abs(row[1] - 1.5) < 0.5       # sd = measured error
    assert str(row[10]).startswith("=")                              # sem derived by formula
    assert abs(row[11] - (row[0] + row[1])) < 1e-9


def test_box_plot_extract_and_export(client):
    """A box plot comes back as kind "box"; the export estimates mean/SD with live formulas."""
    stats = [(2, 5, 7, 9, 13), (4, 8, 11, 15, 19)]
    ok, buf = cv2.imencode(".png", render_box(BoxSpec(stats, ylim=(0, 20))))
    # the bar page finds nothing in a box plot; the box page reads it
    assert _post(client, buf.tobytes(), "box.png").get_json()["figures"] == []
    data = _post(client, buf.tobytes(), "box.png", options=json.dumps({"chart_type": "box"})).get_json()
    panel = data["figures"][0]["panels"][0]
    assert panel["kind"] == "box" and len(panel["boxes"]) == 2
    assert panel["whisker_type"] == "unknown"
    assert abs(panel["boxes"][0]["median"] - 7) < 0.2
    panel["boxes"][0]["n"] = 12
    default_export = client.post("/api/export", json={"figures": data["figures"]})
    default_ws = openpyxl.load_workbook(io.BytesIO(default_export.data)).worksheets[2]
    assert default_ws["T2"].value == "unknown"
    assert "O2" not in default_ws["A2"].value and "S2" not in default_ws["B2"].value
    panel["whisker_type"] = "tukey"
    panel["mean_method"] = "not a method"          # falls back to the default
    r = client.post("/api/export", json={"figures": data["figures"]})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb[wb.sheetnames[2]]
    header = [c.value for c in ws[1]]
    assert header[11] == "mean+sd" and header[14:19] == ["lower end", "Q1", "median", "Q3", "upper end"]
    row = [c.value for c in ws[2]]
    assert row[6] == 12
    assert str(row[0]).startswith('=IF(G2="","",') and "NORMSINV" in str(row[1])
    assert abs(row[16] - 7) < 0.2 and row[19] == "tukey" and row[20] == "luo"
    assert [c.value for c in wb["Summary"][2]][4] == "box"
    combined = [c.value for c in wb["All data"][1]]
    assert combined[-9:-4] == ["lower end", "Q1", "median", "Q3", "upper end"] and combined[-1] == "x"


def test_violin_extract_and_export(client):
    import numpy as np
    rng = np.random.default_rng(1)
    spec = ViolinSpec([rng.normal(10, 2, 14), rng.normal(14, 3, 14)], inner="lines")
    ok, buf = cv2.imencode(".png", render_violin(spec))
    data = _post(client, buf.tobytes(), "violin.png", options=json.dumps({"chart_type": "violin"})).get_json()
    panel = data["figures"][0]["panels"][0]
    assert panel["kind"] == "violin" and len(panel["boxes"]) == 2
    assert abs(panel["boxes"][0]["median"] - spec.stats[0][2]) < 0.3
    panel["boxes"][0]["n"] = 14
    r = client.post("/api/export", json={"figures": data["figures"]})
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb[wb.sheetnames[2]]
    assert "Q1" in [c.value for c in ws[1]] and str(ws.cell(2, 1).value).startswith("=IF(")
    assert [c.value for c in wb["Summary"][2]][4] == "violin"


def test_line_extract_and_export(client):
    spec = LineSpec([0, 7, 14, 21], [[20, 22, 25, 27], [16, 17, 18, 19]], errs=[[1] * 4, [0.8] * 4])
    ok, buf = cv2.imencode(".png", render_line(spec))
    data = _post(client, buf.tobytes(), "line.png", options=json.dumps({"chart_type": "line"})).get_json()
    panel = data["figures"][0]["panels"][0]
    assert panel["kind"] == "line" and len(panel["series"]) == 2
    assert all(len(s["points"]) == 4 for s in panel["series"])
    panel["series"][0]["label"] = "Vehicle"
    panel["series"][0]["points"][0]["n"] = 10
    r = client.post("/api/export", json={"figures": data["figures"]})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb[wb.sheetnames[2]]
    header = [c.value for c in ws[1]]
    assert header[-2:] == ["series", "x"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert len(rows) == 8 and rows[0][5] == "Vehicle" and rows[0][6] == 10
    xs = sorted({round(row[-1]) for row in rows})
    assert xs == [0, 7, 14, 21]
    # malformed series are rejected
    bad = {"figures": [{"panels": [{"kind": "line", "series": [{"points": [{"x_px": 1, "y_px": 1, "mean": "x"}]}]}]}]}
    assert client.post("/api/export", json=bad).status_code == 400


def test_export_excluded_and_empty(client):
    data = _post(client, _png_bytes(), "chart.png").get_json()
    fig = data["figures"][0]
    fig["panels"][0]["excluded"] = True
    assert client.post("/api/export", json={"figures": [fig]}).status_code == 400
    assert client.post("/api/export", json={}).status_code == 400


@pytest.mark.parametrize("bad", [
    {"figures": [{"panels": [{"bars": [{"x1": 0, "x2": 1, "top_px": 1, "cap_px": 1, "mean": "abc"}]}]}]},
    {"figures": [{"panels": ["not a panel"]}]},
    {"figures": [{"panels": [{"bars": [{"x1": 0, "x2": 1, "top_px": 1, "cap_px": 1, "mean": 1.0, "n": "eight"}]}]}]},
    {"figures": "nope"},
    {"figures": [{"panels": [{"kind": "box", "boxes": [{"x1": 0, "x2": 1, "q1_px": 1, "median_px": 1, "q3_px": 1,
                                                         "lo_px": 1, "hi_px": 1, "median": "x"}]}]}]},
])
def test_export_rejects_malformed(client, bad):
    r = client.post("/api/export", json=bad)
    assert r.status_code == 400 and "error" in r.get_json()


def test_old_jobs_cleaned(client, tmp_path):
    old = os.path.join(extract_routes.JOBS_DIR, "deadbeef")
    os.makedirs(old, exist_ok=True)
    open(os.path.join(old, "a.jpg"), "wb").close()
    os.utime(old, (0, 0))
    _post(client, _png_bytes(), "chart.png")
    assert not os.path.exists(old)
